"""Projects (jobs) and their site visits: create a project, add numbered
visits, carry plans forward, track action items across visits, and check
which previous items each visit's report prints."""

import io
import zipfile

import pymupdf
import pytest
from docx import Document

from webapp import server


@pytest.fixture
def client(tmp_path, monkeypatch):
    data, uploads = tmp_path / "data", tmp_path / "uploads"
    data.mkdir()
    uploads.mkdir()
    monkeypatch.setattr(server, "DATA_DIR", str(data))
    monkeypatch.setattr(server, "UPLOADS", str(uploads))
    monkeypatch.setattr(server, "PROJECTS_FILE", str(data / "projects.json"))
    monkeypatch.setattr(server, "LEGACY_DATA_FILE", str(data / "visit.json"))
    c = server.app.test_client()
    assert c.post("/login", data={"pin": server.PIN}).status_code == 302
    return c


def _job(client, **fields):
    body = {"name": "Elm St Retrofit", "number": "24-117", "address": "1 Elm St"}
    body.update(fields)
    r = client.post("/api/jobs", json=body)
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()


def _visit(client, jid, no, **fields):
    r = client.post(f"/api/jobs/{jid}/visits", json={"visit_no": no, **fields})
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()


def _entry(client, **details):
    item = client.post("/api/item").get_json()
    e = item["photos"][0]
    r = client.patch(f"/api/item/{item['item_id']}/photo/{e['photo_id']}", json=details)
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()


def _text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def _report(client):
    r = client.get("/api/export/docx")
    assert r.status_code == 200
    return _text(Document(io.BytesIO(r.data)))


def test_projects_sort_by_number_with_unnumbered_last(client):
    _job(client, name="B", number="24-117")
    _job(client, name="A", number="9-001")
    client.post("/api/projects", json={"name": "Old report"})     # no number
    names = [j["name"] for j in client.get("/api/jobs").get_json()["jobs"]]
    assert names == ["A", "B", "Old report"]


def test_project_number_is_required(client):
    r = client.post("/api/jobs", json={"name": "No number"})
    assert r.status_code == 400 and "number" in r.get_json()["error"]


def test_new_visit_defaults_come_from_the_last_one(client):
    jid = _job(client)["job_id"]
    page = client.get(f"/api/jobs/{jid}").get_json()
    assert page["next_visit_no"] == 1 and page["observed_by"] == ""

    v1 = _visit(client, jid, 1, observed_by="S. Farbic", others_present="GC super",
                stage_of_construction="Framing")
    assert v1["name"] == "Site Visit 1" and v1["observed_by"] == "S. Farbic"
    page = client.get(f"/api/jobs/{jid}").get_json()
    assert page["next_visit_no"] == 2 and page["observed_by"] == "S. Farbic"
    assert page["job"]["visits"] == 1

    # the new visit is the open one
    cur = client.get("/api/visits/current").get_json()
    assert cur["visit"]["project_id"] == v1["project_id"] and cur["job"]["job_id"] == jid


def test_entries_number_through_the_visit_and_print_with_the_visit(client):
    jid = _job(client)["job_id"]
    _visit(client, jid, 2)
    first = _entry(client, description="Cracked beam")
    second = _entry(client, description="Missing hold-down", action="Install the HDU.")
    assert (first["number"], second["number"]) == (1, 2)

    text = _report(client)
    assert "ITEM 2.1\tCracked beam" in text and "ITEM 2.2\tMissing hold-down" in text
    assert "Install the HDU." in text


def _plan_pdf():
    doc = pymupdf.open()
    doc.new_page(width=1200, height=800)
    return io.BytesIO(doc.tobytes())


def test_plans_carry_forward_without_pins(client):
    jid = _job(client)["job_id"]
    _visit(client, jid, 1)
    plan = client.post("/api/plan", data={"plan": (_plan_pdf(), "S-201.pdf")},
                       content_type="multipart/form-data").get_json()
    _entry(client, description="x", location={"plan_id": plan["plan_id"], "page": 1, "x": .5, "y": .5})

    v2 = _visit(client, jid, 2)
    visit = client.get("/api/visit").get_json()
    assert [p["name"] for p in visit["plans"]] == ["S-201.pdf"]
    assert visit["plans"][0]["url"].startswith(f"/uploads/{v2['project_id']}/plans/")
    assert visit["items"] == []
    assert client.get(visit["plans"][0]["url"]).status_code == 200


def test_action_item_prints_until_its_close_out_report(client):
    jid = _job(client)["job_id"]
    v1 = _visit(client, jid, 1)
    e = _entry(client, description="Header undersized at line 2",
               action="The Contractor will install new headers.")
    _entry(client, description="Observation only")
    base = f"/api/visits/{v1['project_id']}/entries/{e['photo_id']}"

    # an update aimed at visit 1's own report is refused
    r = client.post(base + "/updates", json={"text": "x", "visit_no": 1})
    assert r.status_code == 400

    # contractor photos arrive before visit 2: update, then close it out in visit 2
    r = client.post(base + "/updates", json={"text": "Headers on order.", "date": "2026-10-20"})
    assert r.status_code == 200 and r.get_json()["updates"][0]["visit_no"] == 2
    r = client.post(base + "/updates", json={"text": "Headers installed per 4/S8.1.",
                                             "date": "2026-10-29", "close": True})
    assert r.get_json()["closed_visit_no"] == 2

    items = client.get(f"/api/jobs/{jid}").get_json()["action_items"]
    assert [a["label"] for a in items] == ["1.1"] and items[0]["closed_visit_no"] == 2

    _visit(client, jid, 2)
    text = _report(client)
    assert "Previous Action Items" in text and "ITEM 1.1" in text
    assert "Update:	10/20/26 – Headers on order." in text
    assert ("Update:	10/29/26 – Headers installed per 4/S8.1. No further action required."
            in text)
    assert "Observation only" not in text

    _visit(client, jid, 3)
    text = _report(client)
    assert "ITEM 1.1" not in text and "Previous Action Items" not in text


def test_open_items_keep_printing_and_reopen_undoes_a_close(client):
    jid = _job(client)["job_id"]
    v1 = _visit(client, jid, 1)
    e = _entry(client, description="Shear wall nailing", action="Renail per 3/S5.0.")
    base = f"/api/visits/{v1['project_id']}/entries/{e['photo_id']}"
    _visit(client, jid, 2)
    _visit(client, jid, 3)
    assert "ITEM 1.1" in _report(client)          # still open in visit 3

    client.post(base + "/updates", json={"text": "Done", "close": True, "visit_no": 3})
    assert "Done. No further action required." in _report(client)   # period added
    r = client.post(base + "/reopen")
    assert r.get_json()["closed_visit_no"] is None
    text = _report(client)
    assert "Update" in text and "No further action required." not in text

    # deleting the final update reopens too
    client.post(base + "/updates", json={"text": "Done again.", "close": True, "visit_no": 3})
    r = client.delete(base + "/updates/1")
    assert r.get_json()["closed_visit_no"] is None


def test_delete_project_removes_its_visits(client):
    jid = _job(client)["job_id"]
    v1 = _visit(client, jid, 1)
    assert client.delete(f"/api/jobs/{jid}").status_code == 200
    assert client.get("/api/jobs").get_json()["jobs"] == []
    assert client.get(f"/api/visits/{v1['project_id']}").status_code == 404


def test_pages_render(client):
    jid = _job(client)["job_id"]
    v1 = _visit(client, jid, 1)
    for path in ("/", "/settings", "/projects", "/projects/new", f"/projects/{jid}",
                 f"/projects/{jid}/visits/new", f"/projects/{jid}/report",
                 f"/visits/{v1['project_id']}/edit", "/capture"):
        assert client.get(path).status_code == 200, path


def test_open_action_items_log_lists_only_unclosed_items(client):
    jid = _job(client)["job_id"]
    assert client.get(f"/api/jobs/{jid}/open-items.docx").status_code == 404   # no visits yet
    v1 = _visit(client, jid, 1)
    open_ = _entry(client, description="Missing hold-down", action="Install an HDU8.")
    done = _entry(client, description="Header undersized", action="Install new headers.")
    _entry(client, description="Observation only")
    client.post(f"/api/visits/{v1['project_id']}/entries/{open_['photo_id']}/updates",
                json={"text": "HDU8 on order.", "date": "2026-10-20"})
    client.post(f"/api/visits/{v1['project_id']}/entries/{done['photo_id']}/updates",
                json={"text": "Installed.", "close": True})

    r = client.get(f"/api/jobs/{jid}/open-items.docx")

    assert r.status_code == 200
    assert 'Open Action Items.docx"' in r.headers["Content-Disposition"]
    text = _text(Document(io.BytesIO(r.data)))
    assert "Open Action Items" in text and "24-117" in text
    assert "Missing hold-down" in text and "Install an HDU8." in text
    assert "10/20/26 – HDU8 on order." in text
    assert "Header undersized" not in text and "Observation only" not in text


def _jpeg(color="red"):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (80, 60), color).save(buf, "JPEG")
    buf.seek(0)
    return buf


def _shoot(client, entry, color="red"):
    item_id = entry["photo_id"].split("_")[0]
    r = client.post(f"/api/item/{item_id}/photo", content_type="multipart/form-data",
                    data={"photo": (_jpeg(color), "IMG.jpg"), "photo_id": entry["photo_id"]})
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()


def test_project_plans_reach_every_visit_and_removing_one_clears_its_pins(client):
    jid = _job(client)["job_id"]
    _visit(client, jid, 1)
    r = client.post(f"/api/jobs/{jid}/plans", content_type="multipart/form-data",
                    data={"plan": (_plan_pdf(), "S-201.pdf")})
    assert r.status_code == 200
    plan_id = r.get_json()["plan"]["plan_id"]
    e = _entry(client, description="x", location={"plan_id": plan_id, "page": 1, "x": .5, "y": .5})
    _visit(client, jid, 2)                               # new visits get the project's plans
    assert [p["plan_id"] for p in client.get("/api/visit").get_json()["plans"]] == [plan_id]

    r = client.delete(f"/api/jobs/{jid}/plans/{plan_id}")
    assert r.status_code == 200 and r.get_json()["plans"] == []
    assert client.get("/api/visit").get_json()["plans"] == []
    v1 = client.get(f"/api/jobs/{jid}").get_json()["visits"][-1]["project_id"]
    client.post(f"/api/projects/{v1}/open")
    entry = client.get("/api/visit").get_json()["items"][0]["photos"][0]
    assert entry["photo_id"] == e["photo_id"] and entry["location"] is None


def test_several_photos_of_one_item_with_retake_and_delete(client):
    jid = _job(client)["job_id"]
    _visit(client, jid, 1)
    e = _entry(client, description="Cracked beam")
    _shoot(client, e, "red")
    e = _shoot(client, e, "blue")
    first, second = e["images"]
    item_id = e["photo_id"].split("_")[0]

    r = client.post(f"/api/item/{item_id}/photo/{e['photo_id']}/image/{first['image_id']}",
                    content_type="multipart/form-data", data={"photo": (_jpeg("green"), "R.jpg")})
    retaken = r.get_json()["images"]
    assert [im["image_id"] for im in retaken] == [first["image_id"], second["image_id"]]
    assert retaken[0]["filename"] != first["filename"]          # replaced in place

    r = client.delete(f"/api/item/{item_id}/photo/{e['photo_id']}/image/{second['image_id']}")
    assert [im["image_id"] for im in r.get_json()["images"]] == [first["image_id"]]


def test_report_asks_about_items_with_several_photos_and_follows_the_answer(client):
    jid = _job(client)["job_id"]
    v1 = _visit(client, jid, 1)
    e = _entry(client, description="Cracked beam")
    for color in ("red", "blue", "green"):
        e = _shoot(client, e, color)
    single = _entry(client, description="One photo only")
    _shoot(client, single)

    choices = client.get(f"/api/visits/{v1['project_id']}/photo-choices").get_json()["items"]
    assert [c["label"] for c in choices] == ["1.1"] and len(choices[0]["images"]) == 3

    keep = [im["image_id"] for im in e["images"][:2]]
    r = client.post("/api/export/build", json={"photos": {e["photo_id"]: {"images": keep,
                                                                          "layout": "stack"}}})
    assert r.status_code == 200, r.get_data(as_text=True)
    report = zipfile.ZipFile(io.BytesIO(client.get(r.get_json()["url"]).data))
    docx = next(n for n in report.namelist() if n.endswith(".docx"))
    doc = Document(io.BytesIO(report.read(docx)))
    pictures = len(list(doc.element.body.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline")))
    assert pictures == 3          # two of ITEM 1.1's three photos, plus ITEM 1.2's one
